import io
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.models.record import ExpenseRecord


def export_records_to_excel(records: List[ExpenseRecord], company_name: str = "") -> io.BytesIO:
    """
    Generates an Excel workbook from ExpenseRecord objects.
    Reuses the styling, column ordering, and width calculation from logic.py in Pyqt6 project.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Export Table"

    header_font = Font(bold=True)
    header_align = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    # 1. Headers matching your Calling_Page_UI sequence
    headers = [
        "No.",
        "Invoice NO",
        "Project Code",
        "explanation",
        "record_date",
        "amount",
        "expense_center",
        "expense_type",
        "company_name",
        "created_by"
    ]

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.alignment = header_align

    # Freeze header row
    ws.freeze_panes = "A2"

    # 2. Data rows
    for row_idx, r in enumerate(records, start=2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=r.invoice_no)
        ws.cell(row=row_idx, column=3, value=r.project_code)
        ws.cell(row=row_idx, column=4, value=r.explanation or "")
        ws.cell(row=row_idx, column=5, value=r.record_date)

        # Number formatting with commas (e.g. 1,500,000)
        amount_cell = ws.cell(row=row_idx, column=6, value=float(r.amount) if r.amount else 0.0)
        amount_cell.number_format = '#,##0'

        ws.cell(row=row_idx, column=7, value=r.expense_center or "")
        ws.cell(row=row_idx, column=8, value=r.expense_type or "")
        ws.cell(row=row_idx, column=9, value=company_name)
        ws.cell(row=row_idx, column=10, value=r.created_by_name or "")

    # 3. Auto-fit column widths (algorithm from logic.py)
    for col in range(1, len(headers) + 1):
        max_length = 0
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None:
                max_length = max(max_length, len(str(value)))

        letter = get_column_letter(col)
        if col == 1:
            ws.column_dimensions[letter].width = 8
        else:
            ws.column_dimensions[letter].width = max(15, min(max_length + 2, 60))

    # 4. Save to in-memory byte stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output